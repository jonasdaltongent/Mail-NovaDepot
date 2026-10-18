#!/usr/bin/env python3
"""
maak_werkdocumenten.py
Genereert het werkdocument voor les 11 "De zakelijke e-mail" (NovaDepot, ORLO, Communicatie 1).

  CO1_De-zakelijke-e-mail.docx   het werkdocument dat de leerling INLEVERT (een kopie per leerling)

Geen bronbestanden en geen zip: de les werkt alleen in het werkdocument en in de browser. Geen Gmail: de leerlingen
hebben er met hun schoolaccount geen toegang toe (Jonas, 09-10-2026). De e-mail komt in dit document, in een
venster dat eruitziet als "Nieuw bericht", met alleen Aan en Onderwerp (Cc en Bcc volgen in les 12). Er wordt
niets verstuurd.

Gebruik:  python3 maak_werkdocumenten.py
Vereist:  python-docx, Pillow (huisstijl.py en go-dalton-gent.png staan ernaast)

LET OP bij aanpassen:
 - Alle namen, nummers en adressen zijn verzonnen; de adressen eindigen op .example (RFC 2606).
 - De leerling schrijft VOOR Tom (naam en handtekening van Tom): zo staat er geen eigen naam of gegeven in het
   werkdocument.
 - Rij 1 van deel 1 (het lege onderwerp) is de demo van dia 5, al ingevuld in het grijs. Rij 1 van deel 2 ook.
 - De punten staan per deel (4 · 3 · 13) en samen in de controletabel (lesvoorbereiding § 22). Verander je ze, doe
   het op beide plaatsen.
 - Juiste volgorde van deel 2: B · E · F · D · A · C (README §4).
"""
import os

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import RGBColor

import huisstijl as hs   # de huisstijl van het werkdocument (uit _sjablonen/lespakket/werkdocument/)

HERE = os.path.dirname(os.path.abspath(__file__))
UIT = os.path.join(HERE, "CO1_De-zakelijke-e-mail.docx")

# Het kader met de feiten (lesvoorbereiding § 4)
FEITEN = [
    ("Bestelling:", "2026-0418, zuivel voor de winkels, bij Zuivelhoeve De Melkkan"),
    ("Afgesproken levering:", "donderdag 29 oktober 2026, om 7.30 uur"),
    ("Probleem:", "die dag krijgt de koelcel van NovaDepot onderhoud: er kan geen zuivel binnen"),
    ("Nieuwe datum (voorstel van Tom):", "vrijdag 30 oktober 2026, tussen 7.00 en 10.00 uur"),
    ("Contactpersoon:", "mevrouw Kim Wouters, verkoopdienst, verkoop@demelkkan.example"),
    ("Antwoord nodig:", "tegen dinsdag 27 oktober (dan kan Tom nog een andere oplossing zoeken)"),
]

# Deel 1: de slordige mail van Tom, zoals hij ze schreef (alleen lezen)
TOM_AAN = "verkoop@demelkkan.example"
TOM_REGELS = [
    "hey,",
    "ik ben tom van novadepot. kunnen jullie de levering later doen want onze koelcel is kapot!!! \U0001F605",
    "laat iets weten asap",
    "thx",
    "tom",
]

# Deel 1: wat is er mis? Rij 1 = de demo, al ingevuld.
RIJEN_FOUTEN = [
    ("Het onderwerp is leeg: Kim Wouters ziet niet waarover de mail gaat.",
     "Levering bestelling 2026-0418 verplaatsen naar vrijdag 30 oktober"),
    ("", ""),
    ("", ""),
    ("", ""),
    ("", ""),
]

# Deel 2: de mail van Lotte, in zes stukjes door elkaar
STUKJES = [
    ("A", "Met vriendelijke groeten,"),
    ("B", "Beste meneer Van Acker,"),
    ("C", "Lotte Maes · Personeelsdienst · NovaDepot · toestel 205"),
    ("D", "Kunt u mij tegen woensdag 4 november laten weten of dat past? Alvast bedankt."),
    ("E", "Mijn naam is Lotte Maes, van de personeelsdienst van NovaDepot. Ik mail u over het onderhoud van de "
          "koffiemachine in de refter."),
    ("F", "Het onderhoud staat gepland op vrijdag 6 november om 10 uur (opdracht 2026-077). Die dag is de refter "
          "gesloten voor een opleiding. Kan uw technicus op maandag 9 november om 10 uur komen?"),
]

# Deel 2: welke letter hoort bij welke bouwsteen? Rij 1 = de voorbeeldrij, al ingevuld. Juist: B E F D A C.
RIJEN_BOUWSTENEN = [
    ("1  Aanspreking", "B"),
    ("2  Inleiding: wie ben je, waarover gaat het?", ""),
    ("3  De vraag: wat, wanneer en waarom?", ""),
    ("4  Slot: tegen wanneer, bedanken", ""),
    ("5  Slotgroet", ""),
    ("6  Handtekening", ""),
]

# De criteria van lesvoorbereiding § 22 (samen 20): 4 · 3 · 13
CRITERIA = [
    ("Deel 1: ik vond vier fouten en schreef erbij hoe het beter is.", 4),
    ("Deel 2: ik zette de bouwstenen in de juiste volgorde.", 3),
    ("Deel 3: bij Aan staat het adres, en mijn onderwerp zegt welke bestelling het is en wat Tom vraagt.", 2),
    ("Deel 3: ik begin met *Beste mevrouw Wouters,* en sluit af met een slotgroet zonder afkorting.", 2),
    ("Deel 3: mijn inleiding zegt wie schrijft en waarover.", 1),
    ("Deel 3: mijn vraag heeft het bestelnummer, de afgesproken datum, kort de reden en de nieuwe datum met uur.", 4),
    ("Deel 3: mijn slot zegt tegen wanneer Tom een antwoord wil, en bedankt.", 2),
    ("Deel 3: ik schrijf overal *u*. Geen chattaal, geen hoofdletters om te roepen, geen *!!!*, geen emoji's.", 1),
    ("Deel 3: onder mijn mail staat een handtekening met naam, functie, NovaDepot en toestel 214.", 1),
]


def leesvenster(doc, aan, onderwerp, regels):
    """De slordige mail van Tom: een venster dat de leerling alleen leest (grijze rand, geen lichtblauwe vakken)."""
    t = doc.add_table(rows=4, cols=2)
    t.alignment = hs.WD_TABLE_ALIGNMENT.CENTER
    hs.randen(t, buiten="9AA0A6")
    hs.breedtes(t, [3.2, hs.BREEDTE - 3.2])
    kopcel = t.rows[0].cells[0].merge(t.rows[0].cells[1])
    hs.schaduw(kopcel, "F2F6FC")
    hs.celtekst(kopcel, "Nieuw bericht", vet=True, kleur=RGBColor(0x20, 0x21, 0x24))
    kopcel.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    hs.hoogte(t.rows[0], 0.8)
    for i, (veld, waarde, grijs) in enumerate([("Aan", aan, False), ("Onderwerp", onderwerp, True)], start=1):
        links, rechts = t.rows[i].cells
        hs.celtekst(links, veld, kleur=hs.GRIJS)
        hs.celtekst(rechts, waarde, kleur=hs.GRIJS if grijs else None, cursief=grijs)
        for c in (links, rechts):
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            hs.cel_randen(c, "DADCE0", dikte=6, kanten=("bottom",))
        hs.hoogte(t.rows[i], 0.85)
    body = t.rows[3].cells[0].merge(t.rows[3].cells[1])
    hs.marges(body, boven=0.25, onder=0.2, links=0.3, rechts=0.3)
    body.text = ""
    for k, regel in enumerate(regels):
        p = body.paragraphs[0] if k == 0 else body.add_paragraph()
        p.paragraph_format.space_before = hs.Pt(0)
        p.paragraph_format.space_after = hs.Pt(4)
        for deel_tekst in _met_emoji(regel):
            tekst, emoji = deel_tekst
            r = p.add_run(tekst)
            if emoji:
                hs.stijl(r, "Segoe UI Emoji", None)
    hs.niet_splitsen(t)
    hs.ruimte(doc, 4)
    return t


def _met_emoji(regel):
    """Splitst een regel in stukken gewone tekst en emoji (die krijgen een eigen lettertype)."""
    uit, huidig, emoji = [], "", False
    for ch in regel:
        is_emoji = ord(ch) > 0x1F000
        if is_emoji != emoji and huidig:
            uit.append((huidig, emoji))
            huidig = ""
        emoji = is_emoji
        huidig += ch
    if huidig:
        uit.append((huidig, emoji))
    return uit


def nieuwe_bladzijde(doc):
    """Een bladzijdeskeer vóór een deel met puntenvak: page_break_before in een tabelcel (huisstijl.deel met punten)
    negeert Google Documenten. Een kleine lege alinea die zelf op een nieuwe bladzijde begint, doet het wel."""
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    hs.zonder_ruimte(p)
    p.paragraph_format.line_spacing = hs.Pt(2)
    r = p.add_run("")
    r.font.size = hs.Pt(1)
    return p


def werkdocument():
    """In de huisstijl van 12-10-2026 (huisstijl.py): alleen wat de leerling invult, en de feiten die hij
    daarvoor nodig heeft. Hoe hij werkt, staat op de lespagina."""
    doc = hs.nieuw_document()
    doc.core_properties.title = "CO1 De zakelijke e-mail - werkdocument"
    doc.core_properties.author = "Toegepaste Informatica"
    hs.kopblok(doc, "DE ZAKELIJKE E-MAIL", punten=20)

    hs.kader(doc, "Tom vraagt een leverancier om een levering te verplaatsen", [
        "Tom Peeters, magazijnverantwoordelijke bij NovaDepot (toestel 214), vraagt jou: verbeter mijn mail en "
        "schrijf een nieuwe. Ik lees ze na en verstuur ze zelf.",
    ] + FEITEN)

    hs.deel(doc, 1, "Wat is er mis?", punten=4)
    hs.tekst(doc, "Dit is de mail die Tom schreef. Hij is nog niet verstuurd.", klein=True, cursief=True)
    leesvenster(doc, TOM_AAN, "(leeg)", TOM_REGELS)
    hs.tabel(doc, ["Wat is er mis?", "Zo is het beter"], RIJEN_FOUTEN, [8.5, hs.BREEDTE - 8.5],
             invul=(0, 1), rijhoogte=1.25, voorbeeld_rij=1)

    nieuwe_bladzijde(doc)
    hs.deel(doc, 2, "De bouwstenen", punten=3)
    hs.tekst(doc, "Lotte Maes (personeelsdienst) mailt Koffiebranderij De Bonenbaas. Het onderwerp van haar mail: "
                  "Onderhoud koffiemachine verplaatsen naar maandag 9 november. Haar mail staat hier in zes stukjes, "
                  "door elkaar.")
    hs.tabel(doc, ["", "Een stukje van de mail van Lotte"], STUKJES, [1.2, hs.BREEDTE - 1.2], vet=(0,), rijhoogte=0.9)
    hs.vraag(doc, "Welke letter hoort bij welke bouwsteen?")
    hs.tabel(doc, ["Bouwsteen", "Letter"], RIJEN_BOUWSTENEN, [12.6, hs.BREEDTE - 12.6],
             invul=(1,), rijhoogte=0.95, voorbeeld_rij=1)

    nieuwe_bladzijde(doc)
    hs.deel(doc, 3, "De mail van Tom", punten=13)
    hs.tekst(doc, "Schrijf zoals Tom het zou versturen: u-vorm, geen afkortingen, geen emoji.", klein=True,
             cursief=True)
    hs.mailvenster(doc, [("Aanspreking", 1.0), ("Inleiding", 2.0), ("De vraag", 4.2), ("Slot", 2.0),
                         ("Slotgroet", 1.0), ("Handtekening", 2.0)], velden=("Aan", "Onderwerp"))

    hs.deel(doc, 4, "Nalezen", nieuwe_pagina=True)
    # eerst het codewoord van de zelftest (15-10-2026): de lespagina toont het pas als de drie vragen van stap 5
    # juist zijn; zo ziet Jonas wie de zelftest echt maakte (_afspraken/lespagina.md, De laatste stap)
    hs.vraag(doc, "Het codewoord van de zelftest (stap 5 op de lespagina):")
    hs.antwoordvak(doc, 0.9)
    hs.controletabel(doc, CRITERIA)

    hs.deel(doc, None, "Extra — niet verplicht")
    hs.vraag(doc, "Kim Wouters antwoordt: “Vrijdag 30 oktober past, maar pas om 11 uur.” Schrijf een kort antwoord "
                  "van Tom: bedanken en bevestigen.")
    hs.antwoordvak(doc, 3.0)

    doc.save(UIT)


if __name__ == "__main__":
    werkdocument()
    print("Klaar.")
    print("  werkdocument om in te leveren : " + os.path.basename(UIT))
